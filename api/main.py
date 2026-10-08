import os
import re
import json
from pathlib import Path
from typing import List, Optional, Dict, Any

from fastapi import FastAPI, File, Form, UploadFile, HTTPException, Request
from fastapi.responses import JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv
import requests
from bs4 import BeautifulSoup

# Load environment variables from .env
BASE_DIR = Path(__file__).resolve().parent.parent
load_dotenv(dotenv_path=BASE_DIR / ".env")

DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)
PROFILE_FILE = DATA_DIR / "linkedin_profile.md"
PUBLIC_DIR = BASE_DIR / "public"

app = FastAPI(title="LinkedIn Profile AI Chatbot", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class ChatMessage(BaseModel):
    role: str
    content: str


class ChatRequest(BaseModel):
    message: str
    history: Optional[List[ChatMessage]] = []


def get_openai_client():
    """Lazily initialize OpenAI client to avoid failure on startup if key is missing."""
    if "OPENAI_API_KEY" not in os.environ:
        load_dotenv(dotenv_path=BASE_DIR / ".env")
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key or api_key == "your_openai_api_key_here":
        return None, "OpenAI API key is missing or not configured. Please set OPENAI_API_KEY in the .env file."
    try:
        from openai import OpenAI
        client = OpenAI(api_key=api_key)
        return client, None
    except Exception as exc:
        return None, f"Failed to initialize OpenAI client: {str(exc)}"


def clean_and_structure_markdown(raw_text: str, source_title: Optional[str] = None) -> str:
    """Converts extracted text into clean, structured Markdown preserving all facts without inventing data."""
    lines = [line.strip() for line in raw_text.splitlines()]
    
    # Common LinkedIn profile sections to recognize
    recognized_headers = [
        "summary", "about", "about me", "experience", "work experience",
        "education", "skills", "top skills", "licenses & certifications",
        "certifications", "projects", "honors & awards", "awards",
        "volunteer experience", "languages", "publications", "recommendations",
        "contact", "contact info", "patents", "organizations"
    ]
    
    structured_lines = []
    has_title = False
    
    if source_title:
        structured_lines.append(f"# {source_title.strip()}\n")
        has_title = True

    skip_next = False
    for i, line in enumerate(lines):
        if not line:
            continue
        
        lower_line = line.lower().strip(":# -")
        
        # Check if line matches known LinkedIn heading or is an existing Markdown heading
        if line.startswith("#"):
            structured_lines.append(f"\n{line}\n")
            has_title = True
        elif lower_line in recognized_headers:
            heading = line.title().strip(":# -")
            structured_lines.append(f"\n## {heading}\n")
        elif not has_title and i < 3 and len(line) < 60 and not line.startswith("-"):
            # Likely profile name
            structured_lines.append(f"# {line}\n")
            has_title = True
        else:
            # Bullet point or normal text
            if line.startswith(("-", "*", "•", "–")):
                clean_bullet = re.sub(r"^[-*•–]\s*", "", line)
                structured_lines.append(f"- {clean_bullet}")
            else:
                structured_lines.append(line)

    markdown_text = "\n".join(structured_lines)
    # Deduplicate excessive blank lines
    markdown_text = re.sub(r"\n{3,}", "\n\n", markdown_text).strip()
    return markdown_text


def extract_from_pdf(content_bytes: bytes) -> str:
    import pypdf
    import io
    reader = pypdf.PdfReader(io.BytesIO(content_bytes))
    extracted = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            extracted.append(text)
    return "\n\n".join(extracted)


def extract_from_docx(content_bytes: bytes) -> str:
    import docx
    import io
    doc = docx.Document(io.BytesIO(content_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)


def extract_from_html(html_content: str) -> str:
    from markdownify import markdownify as md
    soup = BeautifulSoup(html_content, "html.parser")
    for script_or_style in soup(["script", "style", "noscript", "svg", "header", "footer", "nav"]):
        script_or_style.decompose()
    clean_html = str(soup)
    return md(clean_html, heading_style="ATX", strip=["img"])


def fetch_and_extract_url(url: str) -> str:
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/124.0.0.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
    }
    
    try:
        response = requests.get(url, headers=headers, timeout=12)
    except requests.exceptions.RequestException as e:
        raise HTTPException(
            status_code=400,
            detail=f"Unable to reach URL: {str(e)}. If this page requires login or blocks requests, please export your profile as a PDF or HTML file and upload it."
        )

    # Detect anti-scraping / authentication barriers (such as LinkedIn login redirects)
    if response.status_code in [401, 403, 429, 999] or "authwall" in response.url or "login" in response.url.lower():
        raise HTTPException(
            status_code=400,
            detail=(
                "Access blocked by the target website (anti-scraping or login required). "
                "Please save or print your LinkedIn profile as a PDF or HTML file, "
                "and upload the file directly instead."
            )
        )

    if response.status_code != 200:
        raise HTTPException(
            status_code=400,
            detail=f"Website returned status code {response.status_code}. Please upload a PDF or document file instead."
        )

    # Check content type
    content_type = response.headers.get("content-type", "").lower()
    if "pdf" in content_type:
        raw_text = extract_from_pdf(response.content)
    elif "text/html" in content_type or "xml" in content_type:
        raw_text = extract_from_html(response.text)
    else:
        raw_text = response.text

    if not raw_text.strip():
        raise HTTPException(
            status_code=400,
            detail="No readable text found on the webpage. Please upload a PDF, DOCX, or text file instead."
        )

    soup = BeautifulSoup(response.text[:2000], "html.parser")
    page_title = soup.title.string.strip() if soup.title and soup.title.string else None
    return clean_and_structure_markdown(raw_text, page_title)


def split_profile_sections(markdown_text: str) -> List[Dict[str, str]]:
    """Splits markdown profile into sections based on # or ## headings."""
    lines = markdown_text.splitlines()
    sections = []
    current_heading = "Header"
    current_content = []

    for line in lines:
        if line.startswith(("# ", "## ")):
            if current_content:
                sections.append({
                    "heading": current_heading,
                    "content": "\n".join(current_content).strip()
                })
            current_heading = line.lstrip("#").strip()
            current_content = [line]
        else:
            current_content.append(line)

    if current_content:
        sections.append({
            "heading": current_heading,
            "content": "\n".join(current_content).strip()
        })
    return sections


def select_relevant_context(profile_text: str, user_query: str) -> str:
    """Uses full profile for normal size, or selects relevant sections for large profiles."""
    # Threshold for full profile (~3500 words or ~16,000 characters)
    if len(profile_text) <= 16000:
        return profile_text

    sections = split_profile_sections(profile_text)
    if not sections:
        return profile_text[:12000]

    # Keyword matching against query
    query_tokens = set(re.findall(r"\b[a-zA-Z0-9]{3,}\b", user_query.lower()))
    scored_sections = []

    for idx, sec in enumerate(sections):
        sec_text = (sec["heading"] + " " + sec["content"]).lower()
        score = 0
        
        # Boost for header / summary sections
        if idx in (0, 1) or any(k in sec["heading"].lower() for k in ["summary", "about", "contact"]):
            score += 2

        for token in query_tokens:
            if token in sec["heading"].lower():
                score += 5
            if token in sec_text:
                score += 1

        scored_sections.append((score, sec))

    # Sort descending by score
    scored_sections.sort(key=lambda x: x[0], reverse=True)
    
    # Pick top sections up to 12,000 characters
    selected = []
    total_len = 0
    for _, sec in scored_sections:
        if total_len + len(sec["content"]) < 12000:
            selected.append(sec["content"])
            total_len += len(sec["content"])

    # Maintain original order
    ordered_selected = [s["content"] for s in sections if s["content"] in selected]
    return "\n\n".join(ordered_selected)


@app.get("/api/profile")
def get_profile():
    """Returns the current LinkedIn profile Markdown and existence status."""
    if not PROFILE_FILE.exists():
        return {
            "exists": False,
            "markdown": "",
            "message": "No profile loaded yet. Please upload a profile file or provide a public URL."
        }
    
    try:
        content = PROFILE_FILE.read_text(encoding="utf-8")
        return {
            "exists": bool(content.strip()),
            "markdown": content,
            "file_size": len(content.encode("utf-8")),
            "last_modified": os.path.getmtime(PROFILE_FILE)
        }
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Error reading profile file: {str(exc)}")


@app.post("/api/upload")
async def upload_profile(
    file: Optional[UploadFile] = File(None),
    url: Optional[str] = Form(None)
):
    """
    Accepts file upload (PDF, DOC, DOCX, TXT, HTML, MD) or a public webpage URL.
    Extracts text locally, converts to structured Markdown, and saves to data/linkedin_profile.md.
    """
    clean_markdown = ""

    if url and url.strip():
        clean_markdown = fetch_and_extract_url(url.strip())

    elif file is not None:
        filename = file.filename or "unknown"
        ext = Path(filename).suffix.lower()
        allowed_extensions = {".pdf", ".docx", ".doc", ".txt", ".md", ".html", ".htm"}
        
        if ext not in allowed_extensions:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file format '{ext}'. Allowed: PDF, DOCX, DOC, TXT, HTML, MD"
            )

        contents = await file.read()
        
        # 15MB limit check
        if len(contents) > 15 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="File size exceeds maximum limit of 15MB.")

        if not contents:
            raise HTTPException(status_code=400, detail="Uploaded file is empty.")

        try:
            if ext == ".pdf":
                raw_text = extract_from_pdf(contents)
            elif ext == ".docx":
                raw_text = extract_from_docx(contents)
            elif ext == ".doc":
                # For legacy .doc, try decoding or text fallback
                try:
                    raw_text = contents.decode("utf-8", errors="ignore")
                    # Clean binary artifacts
                    raw_text = "".join(ch for ch in raw_text if ch.isprintable() or ch in "\n\r\t")
                except Exception:
                    raise HTTPException(
                        status_code=400,
                        detail="Cannot parse legacy .doc binary format directly. Please save as .docx or PDF and upload."
                    )
            elif ext in (".html", ".htm"):
                raw_text = extract_from_html(contents.decode("utf-8", errors="replace"))
            else:  # .txt, .md
                raw_text = contents.decode("utf-8", errors="replace")

            if not raw_text.strip():
                raise HTTPException(status_code=400, detail="Could not extract any readable text from the file.")

            clean_markdown = clean_and_structure_markdown(raw_text, Path(filename).stem)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"Failed to process document: {str(exc)}")

    else:
        raise HTTPException(status_code=400, detail="Please provide either a file upload or a webpage URL.")

    # Save to data/linkedin_profile.md
    try:
        PROFILE_FILE.write_text(clean_markdown, encoding="utf-8")
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to save profile markdown: {str(exc)}")

    return {
        "status": "success",
        "message": "Profile successfully processed and updated.",
        "markdown": clean_markdown
    }


@app.post("/api/chat")
async def chat(request: ChatRequest):
    """
    Answers questions about the profile using data/linkedin_profile.md as the sole factual source.
    """
    if not PROFILE_FILE.exists() or not PROFILE_FILE.read_text(encoding="utf-8").strip():
        return {
            "reply": "Please upload or import a LinkedIn profile first so I can answer questions about it.",
            "section": "None"
        }

    user_query = request.message.strip()
    if not user_query:
        return {
            "reply": "Please enter a question.",
            "section": "None"
        }

    client, err_msg = get_openai_client()
    if err_msg or not client:
        return {
            "reply": f"⚠️ {err_msg}",
            "section": "Configuration Required"
        }

    profile_text = PROFILE_FILE.read_text(encoding="utf-8")
    context_markdown = select_relevant_context(profile_text, user_query)

    model_name = os.getenv("OPENAI_MODEL", "gpt-4o-mini").strip() or "gpt-4o-mini"

    system_prompt = (
        "You are the individual described in the LinkedIn profile below. "
        "Answer all questions about yourself in the first person ('I', 'me', 'my').\n\n"
        "STRICT FACTUAL GUIDELINES:\n"
        "1. Base your answer SOLELY and STRICTLY on the facts provided in the Profile Context.\n"
        "2. Do NOT extrapolate, invent, hallucinate, or assume any facts not stated in the profile.\n"
        "3. If the profile does not provide enough information to answer the question, you MUST reply with this exact sentence:\n"
        "   \"My profile does not include that information.\"\n"
        "4. Keep your answer natural, concise, and direct.\n"
        "5. Indicate the exact Markdown section heading from the profile that directly supports your answer "
        "(e.g., 'Summary', 'Experience', 'Education', 'Skills', 'Certifications & Achievements', 'Projects', etc.). "
        "If information is missing or not found, set the section heading to 'None'.\n\n"
        "RESPONSE FORMAT:\n"
        "You must respond with valid JSON containing exactly these two keys:\n"
        "{\n"
        "  \"reply\": \"<your concise first-person answer here>\",\n"
        "  \"section\": \"<supporting section heading or None>\"\n"
        "}\n\n"
        f"--- PROFILE CONTEXT ---\n{context_markdown}\n--- END PROFILE CONTEXT ---"
    )

    messages = [{"role": "system", "content": system_prompt}]

    # Include recent chat history (up to last 6 turns)
    if request.history:
        for turn in request.history[-6:]:
            if turn.role in ("user", "assistant"):
                messages.append({"role": turn.role, "content": turn.content})

    messages.append({"role": "user", "content": user_query})

    try:
        response = client.chat.completions.create(
            model=model_name,
            messages=messages,
            temperature=0.1,
            response_format={"type": "json_object"}
        )
        raw_content = response.choices[0].message.content
        parsed = json.loads(raw_content)
        
        reply = parsed.get("reply", "My profile does not include that information.")
        section = parsed.get("section", "None")
        
        return {
            "reply": reply,
            "section": section
        }

    except Exception as exc:
        return {
            "reply": f"An error occurred while contacting OpenAI: {str(exc)}",
            "section": "Error"
        }


# Mount static assets and serve index.html
if PUBLIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(PUBLIC_DIR)), name="static")

    @app.get("/")
    async def serve_index():
        index_file = PUBLIC_DIR / "index.html"
        if index_file.exists():
            return FileResponse(index_file)
        return JSONResponse({"status": "healthy", "service": "LinkedIn Profile Chatbot"})
